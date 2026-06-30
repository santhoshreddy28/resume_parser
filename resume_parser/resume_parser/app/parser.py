"""
Smart Resume Parser - core parsing engine.

Extracts structured information (contact details, skills, education,
experience) from PDF / DOCX resumes using PyMuPDF / python-docx for text
extraction and spaCy (rule-based PhraseMatcher) + regex for information
extraction.
"""

import re
import json
from pathlib import Path

import fitz  # PyMuPDF
import docx
import spacy
from spacy.matcher import PhraseMatcher

# --------------------------------------------------------------------------
# 1. TEXT EXTRACTION
# --------------------------------------------------------------------------

def extract_text_from_pdf(file_path: str) -> str:
    """Extract raw text from a PDF resume using PyMuPDF."""
    text_parts = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    """Extract raw text from a DOCX resume using python-docx."""
    document = docx.Document(file_path)
    lines = [p.text for p in document.paragraphs]
    # Also pull text out of any tables (common in resume templates)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)
    return "\n".join(lines)


def extract_text(file_path: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


# --------------------------------------------------------------------------
# 2. CLEANING / PREPROCESSING
# --------------------------------------------------------------------------

def clean_text(text: str) -> str:
    """Normalise whitespace and drop empty lines."""
    text = text.replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# 3. CONTACT INFO (regex)
# --------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(\+?\d{1,3}[-.\s]?)?(\(?\d{3,5}\)?[-.\s.]?){1,3}\d{3,5}")
LINKEDIN_RE = re.compile(r"(https?://)?(www\.)?linkedin\.com/[A-Za-z0-9_\-/]+", re.I)
GITHUB_RE = re.compile(r"(https?://)?(www\.)?github\.com/[A-Za-z0-9_\-/]+", re.I)


def extract_email(text: str):
    m = EMAIL_RE.search(text)
    return m.group(0) if m else None


def extract_phone(text: str):
    for m in PHONE_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group(0))
        if 9 <= len(digits) <= 13:
            return m.group(0).strip()
    return None


def extract_linkedin(text: str):
    m = LINKEDIN_RE.search(text)
    return m.group(0) if m else None


def extract_github(text: str):
    m = GITHUB_RE.search(text)
    return m.group(0) if m else None


def extract_name(text: str):
    """Heuristic: the first line that looks like a person's name
    (1-4 capitalised words, no digits/@/section keywords)."""
    section_words = {"resume", "curriculum", "vitae", "cv", "profile"}
    for line in text.split("\n")[:6]:
        candidate = line.strip()
        if not candidate or any(c.isdigit() for c in candidate):
            continue
        if "@" in candidate or "http" in candidate.lower():
            continue
        words = candidate.split()
        if 1 <= len(words) <= 4 and all(w[0].isupper() for w in words if w[0].isalpha()):
            if candidate.lower() not in section_words:
                return candidate
    return None


# --------------------------------------------------------------------------
# 4. SECTION SPLITTING
# --------------------------------------------------------------------------

SECTION_HEADERS = {
    "summary": ["summary", "objective", "profile"],
    "skills": ["skills", "technical skills", "core competencies", "key skills"],
    "experience": ["experience", "work experience", "professional experience",
                   "employment history", "work history"],
    "education": ["education", "academic background", "academics"],
    "projects": ["projects", "academic projects", "personal projects"],
    "certifications": ["certifications", "certificates", "licenses"],
}

# Build a reverse lookup: header text -> canonical section name
HEADER_LOOKUP = {}
for canonical, variants in SECTION_HEADERS.items():
    for v in variants:
        HEADER_LOOKUP[v] = canonical


def split_into_sections(text: str) -> dict:
    """Split resume text into labelled sections based on heading lines."""
    lines = text.split("\n")
    sections = {"header": []}
    current = "header"

    for line in lines:
        stripped = line.strip()
        key = stripped.lower().strip(":")
        # A heading is short, matches a known keyword, and isn't a sentence
        if key in HEADER_LOOKUP and len(stripped) <= 40:
            current = HEADER_LOOKUP[key]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(stripped)

    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


# --------------------------------------------------------------------------
# 5. SKILL EXTRACTION (spaCy PhraseMatcher - rule based, no internet model needed)
# --------------------------------------------------------------------------

SKILL_LIBRARY = [
    # Languages
    "Python", "Java", "C++", "C", "C#", "JavaScript", "TypeScript", "R", "Go", "Rust",
    "PHP", "Ruby", "Swift", "Kotlin", "SQL", "Scala", "MATLAB",
    # Web
    "HTML", "CSS", "React", "Angular", "Vue.js", "Node.js", "Django", "Flask",
    "FastAPI", "Express.js", "Bootstrap", "REST API", "GraphQL",
    # Data / ML
    "Pandas", "NumPy", "Scikit-learn", "TensorFlow", "PyTorch", "Keras",
    "Machine Learning", "Deep Learning", "NLP", "Computer Vision", "spaCy",
    "OpenCV", "Matplotlib", "Seaborn", "Data Analysis", "Data Visualization",
    "Power BI", "Tableau", "Excel",
    # Databases
    "MySQL", "PostgreSQL", "MongoDB", "SQLite", "Oracle", "Redis", "Firebase",
    # Cloud / DevOps
    "AWS", "Azure", "Google Cloud", "Docker", "Kubernetes", "Git", "GitHub",
    "CI/CD", "Jenkins", "Linux",
    # Tools / Other
    "Streamlit", "Tkinter", "Selenium", "BeautifulSoup", "Postman",
    "Jupyter", "VS Code", "Agile", "Scrum", "JIRA",
]

_nlp = spacy.blank("en")
_matcher = PhraseMatcher(_nlp.vocab, attr="LOWER")
_matcher.add("SKILLS", [_nlp.make_doc(skill) for skill in SKILL_LIBRARY])


def extract_skills(text: str) -> list:
    doc = _nlp.make_doc(text)
    matches = _matcher(doc)
    found = set()
    for match_id, start, end in matches:
        found.add(doc[start:end].text)
    # Normalise casing back to the canonical entry in SKILL_LIBRARY
    canonical = {s.lower(): s for s in SKILL_LIBRARY}
    return sorted({canonical.get(f.lower(), f) for f in found})


# --------------------------------------------------------------------------
# 6. EDUCATION / EXPERIENCE EXTRACTION (regex + line heuristics)
# --------------------------------------------------------------------------

DEGREE_KEYWORDS = [
    "b.tech", "btech", "bachelor", "b.sc", "bsc", "b.e", "be ", "m.tech", "mtech",
    "master", "m.sc", "msc", "mba", "phd", "ph.d", "diploma", "associate degree",
]

DATE_RANGE_RE = re.compile(
    r"((19|20)\d{2})\s*(-|to|–|—)\s*((19|20)\d{2}|present|current)", re.I
)


def extract_education(section_text: str) -> list:
    entries = []
    if not section_text:
        return entries
    for line in section_text.split("\n"):
        low = line.lower()
        if any(k in low for k in DEGREE_KEYWORDS) or DATE_RANGE_RE.search(line):
            date_match = DATE_RANGE_RE.search(line)
            entries.append({
                "line": line,
                "years": date_match.group(0) if date_match else None,
            })
    return entries


def extract_experience(section_text: str) -> list:
    entries = []
    if not section_text:
        return entries
    current_entry = None
    for line in section_text.split("\n"):
        date_match = DATE_RANGE_RE.search(line)
        if date_match or re.match(r"^[A-Z][A-Za-z0-9&.,\-\s]{2,60}(\||-|–)\s*[A-Z]", line):
            if current_entry:
                entries.append(current_entry)
            current_entry = {
                "title_line": line,
                "years": date_match.group(0) if date_match else None,
                "bullets": [],
            }
        elif current_entry is not None:
            current_entry["bullets"].append(line)
    if current_entry:
        entries.append(current_entry)
    return entries


# --------------------------------------------------------------------------
# 7. MAIN ORCHESTRATOR
# --------------------------------------------------------------------------

def parse_resume(file_path: str) -> dict:
    raw_text = extract_text(file_path)
    text = clean_text(raw_text)
    sections = split_into_sections(text)

    result = {
        "file_name": Path(file_path).name,
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "linkedin": extract_linkedin(text),
        "github": extract_github(text),
        "skills": extract_skills(sections.get("skills", "") or text),
        "education": extract_education(sections.get("education", "")),
        "experience": extract_experience(sections.get("experience", "")),
        "summary": sections.get("summary", None),
        "sections_detected": list(sections.keys()),
    }
    return result


def parse_resume_from_text(text: str, file_name: str = "uploaded_resume") -> dict:
    """Same as parse_resume but starting from raw text (used by the
    Streamlit UI when reading an uploaded file object directly)."""
    text = clean_text(text)
    sections = split_into_sections(text)
    result = {
        "file_name": file_name,
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "linkedin": extract_linkedin(text),
        "github": extract_github(text),
        "skills": extract_skills(sections.get("skills", "") or text),
        "education": extract_education(sections.get("education", "")),
        "experience": extract_experience(sections.get("experience", "")),
        "summary": sections.get("summary", None),
        "sections_detected": list(sections.keys()),
    }
    return result


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python parser.py <resume_file>")
        sys.exit(1)
    data = parse_resume(sys.argv[1])
    print(json.dumps(data, indent=2))
