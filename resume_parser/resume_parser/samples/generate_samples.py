"""Generate 5 realistic sample resumes (mix of PDF and DOCX) for testing
the Smart Resume Parser."""

from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

OUT_DIR = Path("/home/claude/resume_parser/samples")
OUT_DIR.mkdir(parents=True, exist_ok=True)

styles = getSampleStyleSheet()
h_style = ParagraphStyle("Heading", parent=styles["Heading2"], textColor=colors.HexColor("#1a1a1a"))
name_style = ParagraphStyle("Name", parent=styles["Title"], fontSize=18)
body_style = styles["Normal"]


def make_pdf_resume(filename, name, contact_line, summary, skills, experience, education):
    path = OUT_DIR / filename
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    story = [Paragraph(name, name_style), Paragraph(contact_line, body_style), Spacer(1, 12)]

    story.append(Paragraph("Summary", h_style))
    story.append(Paragraph(summary, body_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Skills", h_style))
    story.append(Paragraph(", ".join(skills), body_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Experience", h_style))
    for exp in experience:
        story.append(Paragraph(exp["title"], body_style))
        for b in exp["bullets"]:
            story.append(Paragraph(f"- {b}", body_style))
        story.append(Spacer(1, 6))
    story.append(Spacer(1, 6))

    story.append(Paragraph("Education", h_style))
    for edu in education:
        story.append(Paragraph(edu, body_style))

    doc.build(story)
    print(f"Created {path}")


def make_docx_resume(filename, name, contact_line, summary, skills, experience, education):
    path = OUT_DIR / filename
    d = docx.Document()
    d.add_heading(name, level=0)
    d.add_paragraph(contact_line)

    d.add_heading("Summary", level=1)
    d.add_paragraph(summary)

    d.add_heading("Skills", level=1)
    d.add_paragraph(", ".join(skills))

    d.add_heading("Experience", level=1)
    for exp in experience:
        d.add_paragraph(exp["title"])
        for b in exp["bullets"]:
            d.add_paragraph(b, style="List Bullet")

    d.add_heading("Education", level=1)
    for edu in education:
        d.add_paragraph(edu)

    d.save(str(path))
    print(f"Created {path}")


# ---------------------------------------------------------------------
# Sample 1 - PDF - Data/ML resume
make_pdf_resume(
    "sample_1_priya_sharma.pdf",
    "Priya Sharma",
    "priya.sharma92@gmail.com | +91 98765 43210 | linkedin.com/in/priyasharma | github.com/priyasharma",
    "Final-year Computer Science student with hands-on experience in machine learning "
    "and data analysis, seeking an internship to apply ML and Python skills to real-world problems.",
    ["Python", "Pandas", "NumPy", "Scikit-learn", "TensorFlow", "SQL", "Matplotlib", "Git"],
    [
        {"title": "Machine Learning Intern - DataWiz Analytics | 2025 - 2026",
         "bullets": [
             "Built a classification model improving prediction accuracy by 12%",
             "Cleaned and preprocessed datasets of over 100,000 rows using Pandas",
         ]},
        {"title": "Research Assistant - University AI Lab | 2024 - 2025",
         "bullets": ["Assisted in NLP research using spaCy and TensorFlow"]},
    ],
    ["B.Tech in Computer Science, Punjab University, 2022 - 2026"],
)

# ---------------------------------------------------------------------
# Sample 2 - DOCX - Web Development resume
make_docx_resume(
    "sample_2_arjun_mehta.docx",
    "Arjun Mehta",
    "arjun.mehta@outlook.com | +91 91234 56789 | github.com/arjunmehta",
    "Full-stack web developer experienced in building responsive applications using "
    "React, Node.js and REST APIs.",
    ["JavaScript", "React", "Node.js", "Express.js", "MongoDB", "HTML", "CSS", "Docker"],
    [
        {"title": "Full Stack Developer Intern - Webify Solutions | 2025 - 2026",
         "bullets": [
             "Developed REST API endpoints used by over 5,000 daily users",
             "Implemented React components reducing page load time by 20%",
         ]},
    ],
    ["B.Sc in Information Technology, Delhi University, 2021 - 2025"],
)

# ---------------------------------------------------------------------
# Sample 3 - PDF - Cybersecurity resume
make_pdf_resume(
    "sample_3_neha_kapoor.pdf",
    "Neha Kapoor",
    "neha.kapoor@protonmail.com | +91 99887 76655 | linkedin.com/in/nehakapoor",
    "Cybersecurity enthusiast with strong fundamentals in network security, Linux "
    "administration, and penetration testing.",
    ["Python", "Linux", "Selenium", "SQL", "Git", "AWS", "Postman"],
    [
        {"title": "Security Intern - CyberShield Pvt Ltd | 2024 - 2025",
         "bullets": ["Performed vulnerability assessments on 10+ internal applications"]},
    ],
    ["B.Tech in Electronics and Communication, NIT Jalandhar, 2021 - 2025",
     "Certified Ethical Hacker (CEH), 2025"],
)

# ---------------------------------------------------------------------
# Sample 4 - DOCX - Business/Analytics resume
make_docx_resume(
    "sample_4_rohan_verma.docx",
    "Rohan Verma",
    "rohan.verma@gmail.com | +91 90909 80808",
    "Business analyst with experience in Excel-based reporting, dashboard design, "
    "and stakeholder communication.",
    ["Excel", "Power BI", "Tableau", "SQL", "Python", "Pandas"],
    [
        {"title": "Business Analyst Intern - FinEdge Corp | 2025",
         "bullets": [
             "Built Power BI dashboards tracking quarterly revenue trends",
             "Automated monthly reporting using Python and Pandas",
         ]},
    ],
    ["MBA in Finance, Amritsar Business School, 2023 - 2025",
     "B.Com, Guru Nanak Dev University, 2019 - 2022"],
)

# ---------------------------------------------------------------------
# Sample 5 - PDF - Entry-level / minimal resume (stress test)
make_pdf_resume(
    "sample_5_simran_kaur.pdf",
    "Simran Kaur",
    "simran.kaur@yahoo.com | +91 98123 45670",
    "Recent graduate eager to start a career in software development.",
    ["Java", "C++", "MySQL", "Git"],
    [],
    ["B.Tech in Computer Science, Thapar Institute, 2022 - 2026"],
)

print("\nAll 5 sample resumes generated.")
